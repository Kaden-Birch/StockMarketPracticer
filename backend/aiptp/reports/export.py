"""Report exports: CSV, JSON, Markdown, XLSX, and PDF — including AI
recommendation history (PRD §26)."""

import csv
import io
import json
from datetime import datetime, timezone
from typing import Any

TXN_COLUMNS = [
    "executed_at", "symbol", "kind", "side", "quantity", "price",
    "amount", "realized_pnl", "fx_rate", "quote_currency", "origin",
]


def _txn_rows(transactions) -> list[dict[str, Any]]:
    rows = []
    for t in transactions:
        rows.append(
            {
                "executed_at": t.executed_at.isoformat(),
                "symbol": t.symbol,
                "kind": t.kind.value,
                "side": t.side.value,
                "quantity": str(t.quantity),
                "price": str(t.price),
                "amount": str(t.amount),
                "realized_pnl": str(t.realized_pnl) if t.realized_pnl is not None else "",
                "fx_rate": str(t.fx_rate),
                "quote_currency": t.quote_currency,
                "origin": t.origin.value,
            }
        )
    return rows


def to_csv(portfolio_view: dict[str, Any], transactions) -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=TXN_COLUMNS)
    writer.writeheader()
    writer.writerows(_txn_rows(transactions))
    return buf.getvalue()


def to_json(portfolio_view: dict[str, Any], transactions, analytics: dict | None) -> str:
    return json.dumps(
        {
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "portfolio": portfolio_view,
            "transactions": _txn_rows(transactions),
            "analytics": analytics,
        },
        indent=2,
        default=str,
    )


def _rec_rows(recommendations) -> list[dict[str, Any]]:
    rows = []
    for r in recommendations or []:
        rows.append(
            {
                "created_at": r.created_at.isoformat(),
                "model_id": r.model_id,
                "action": r.action.value,
                "symbol": r.symbol,
                "sizing": r.sizing,
                "confidence": str(r.confidence) if r.confidence is not None else "",
                "status": r.status.value,
                "rationale": (r.rationale or "")[:300],
            }
        )
    return rows


def to_xlsx(portfolio_view: dict[str, Any], transactions, analytics: dict | None,
            recommendations=None) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    p = portfolio_view

    summary = wb.active
    summary.title = "Summary"
    summary.append(["AIPTP portfolio report", p["name"]])
    summary.append(["Generated (UTC)", datetime.now(timezone.utc).isoformat()])
    summary.append(["Paper trading only — simulated funds, real market data."])
    summary.append([])
    for label, key in [
        ("Currency", "currency"), ("Total value", "total_value"),
        ("Cash", "cash_balance"), ("Market value", "market_value"),
        ("Unrealized P&L", "unrealized_pnl"), ("Lifetime return", "lifetime_return"),
        ("Cost basis method", "cost_basis_method"),
    ]:
        summary.append([label, p[key]])
    if analytics and analytics.get("risk"):
        summary.append([])
        summary.append(["Risk"])
        for k, v in analytics["risk"].items():
            summary.append([k, v])

    holdings = wb.create_sheet("Holdings")
    holdings.append(["Symbol", "Shares", "Avg cost", "Price", "Market value", "Unrealized"])
    for h in p["holdings"]:
        holdings.append([h["symbol"], h["quantity"], h["avg_cost"], h["price"],
                         h["market_value"], h["unrealized_pnl"]])

    txns = wb.create_sheet("Transactions")
    txns.append([c.replace("_", " ").title() for c in TXN_COLUMNS])
    for row in _txn_rows(transactions):
        txns.append([row[c] for c in TXN_COLUMNS])

    recs = wb.create_sheet("AI Recommendations")
    rec_cols = ["created_at", "model_id", "action", "symbol", "sizing",
                "confidence", "status", "rationale"]
    recs.append([c.replace("_", " ").title() for c in rec_cols])
    for row in _rec_rows(recommendations):
        recs.append([row[c] for c in rec_cols])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(portfolio_view: dict[str, Any], transactions, analytics: dict | None,
           recommendations=None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    p = portfolio_view
    styles = getSampleStyleSheet()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=14 * mm, rightMargin=14 * mm)
    grid = TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef7")),
    ])
    story = [
        Paragraph(f"AIPTP portfolio report — {p['name']}", styles["Title"]),
        Paragraph(
            f"Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} · "
            f"currency {p['currency']} · cost basis {p['cost_basis_method']}",
            styles["Normal"],
        ),
        Paragraph("Paper trading only — simulated funds, real market data. "
                  "AI output is informational and educational, not financial advice.",
                  styles["Italic"]),
        Spacer(1, 6 * mm),
        Table([
            ["Total value", p["total_value"], "Cash", p["cash_balance"],
             "Unrealized P&L", p["unrealized_pnl"], "Lifetime", p["lifetime_return"]],
        ], style=grid),
        Spacer(1, 5 * mm),
        Paragraph("Holdings", styles["Heading2"]),
        Table(
            [["Symbol", "Shares", "Avg cost", "Price", "Market value", "Unrealized"]]
            + [[h["symbol"], h["quantity"], h["avg_cost"] or "—", h["price"] or "—",
                h["market_value"] or "—", h["unrealized_pnl"] or "—"]
               for h in p["holdings"]],
            style=grid, repeatRows=1,
        ),
    ]
    if analytics and analytics.get("risk"):
        risk = analytics["risk"]
        story += [
            Spacer(1, 5 * mm),
            Paragraph("Risk", styles["Heading2"]),
            Table([["Sharpe", "Sortino", "Beta", "Volatility", "Max drawdown"],
                   [risk.get("sharpe"), risk.get("sortino"), risk.get("beta"),
                    risk.get("volatility"), risk.get("max_drawdown")]], style=grid),
        ]
    rec_rows = _rec_rows(recommendations)
    if rec_rows:
        story += [
            Spacer(1, 5 * mm),
            Paragraph("AI recommendation history", styles["Heading2"]),
            Table(
                [["Created", "Model", "Action", "Symbol", "Confidence", "Status"]]
                + [[r["created_at"][:16], r["model_id"], r["action"], r["symbol"],
                    r["confidence"], r["status"]] for r in rec_rows[:60]],
                style=grid, repeatRows=1,
            ),
        ]
    story += [
        Spacer(1, 5 * mm),
        Paragraph("Transactions", styles["Heading2"]),
        Table(
            [["Executed", "Symbol", "Kind", "Side", "Qty", "Price", "Amount", "Origin"]]
            + [[r["executed_at"][:16], r["symbol"], r["kind"], r["side"], r["quantity"],
                r["price"], r["amount"], r["origin"]] for r in _txn_rows(transactions)[:200]],
            style=grid, repeatRows=1,
        ),
    ]
    doc.build(story)
    return buf.getvalue()


def to_markdown(portfolio_view: dict[str, Any], transactions, analytics: dict | None) -> str:
    p = portfolio_view
    lines = [
        f"# Portfolio report — {p['name']}",
        "",
        f"Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} · "
        f"currency {p['currency']} · cost basis {p['cost_basis_method']}",
        "",
        "## Summary",
        "",
        f"| Metric | Value |",
        f"| --- | ---: |",
        f"| Total value | {p['total_value']} |",
        f"| Cash | {p['cash_balance']} |",
        f"| Market value | {p['market_value']} |",
        f"| Unrealized P&L | {p['unrealized_pnl']} |",
        f"| Lifetime return | {p['lifetime_return']} |",
        "",
        "## Holdings",
        "",
        "| Symbol | Shares | Avg cost | Price | Market value | Unrealized |",
        "| --- | ---: | ---: | ---: | ---: | ---: |",
    ]
    for h in p["holdings"]:
        lines.append(
            f"| {h['symbol']} | {h['quantity']} | {h['avg_cost'] or '—'} | "
            f"{h['price'] or '—'} | {h['market_value'] or '—'} | {h['unrealized_pnl'] or '—'} |"
        )
    if analytics:
        risk = analytics.get("risk", {})
        records = analytics.get("records", {})
        lines += [
            "",
            "## Analytics",
            "",
            "| Metric | Value |",
            "| --- | ---: |",
            f"| Sharpe | {risk.get('sharpe', '—')} |",
            f"| Sortino | {risk.get('sortino', '—')} |",
            f"| Beta | {risk.get('beta', '—')} |",
            f"| Volatility (ann.) | {risk.get('volatility', '—')}% |",
            f"| Max drawdown | {risk.get('max_drawdown', '—')}% |",
            f"| Win rate | {records.get('win_rate', '—')}% |",
            f"| Avg holding period | {records.get('avg_holding_days', '—')} days |",
            f"| Diversification score | {analytics.get('diversification', {}).get('score', '—')} |",
        ]
    lines += ["", "## Transactions", "",
              "| Executed | Symbol | Kind | Side | Qty | Price | Amount | Realized | Origin |",
              "| --- | --- | --- | --- | ---: | ---: | ---: | ---: | --- |"]
    for r in _txn_rows(transactions):
        lines.append(
            f"| {r['executed_at'][:16]} | {r['symbol']} | {r['kind']} | {r['side']} | "
            f"{r['quantity']} | {r['price']} | {r['amount']} | {r['realized_pnl'] or '—'} | {r['origin']} |"
        )
    lines.append("")
    lines.append("*Paper trading only — simulated funds, real market data.*")
    return "\n".join(lines)
